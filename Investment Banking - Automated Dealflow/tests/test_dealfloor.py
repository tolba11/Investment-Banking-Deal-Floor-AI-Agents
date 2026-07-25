"""No network, no keys. Providers are stubbed at the seam."""

from __future__ import annotations

import io

import pytest

from dealfloor import catalog, chat, intake, orchestrator, providers
from dealfloor.orchestrator import Run, RunConfig


# ── stub ─────────────────────────────────────────────────────────────────

@pytest.fixture
def stub(monkeypatch):
    """Replace the one function that talks to the outside world."""
    calls: list[dict] = []

    def fake(model_key, system, content, *, max_tokens=1400,
             web_search=False, history=()):
        calls.append({
            "model": model_key, "system": system,
            "content": content, "web_search": web_search,
        })
        model = providers.BY_KEY[model_key]
        return providers.Reply(
            text="## Output\n- Revenue AED 15.2bn\n- EV/EBITDA 6.4x",
            model=model.label, vendor=model.vendor,
            citations=["https://example.com/ir"] if web_search else [],
            input_tokens=800, output_tokens=250,
        )

    monkeypatch.setattr(providers, "complete", fake)
    monkeypatch.setattr(orchestrator.providers, "complete", fake)
    monkeypatch.setattr(intake.providers, "complete", fake)
    monkeypatch.setattr(chat.providers, "complete", fake)
    return calls


# ── catalogue ────────────────────────────────────────────────────────────

def test_all_42_agents_load():
    agents = catalog.agents()
    assert len(agents) == 42
    for a in agents.values():
        assert a.title and a.deliverable and a.steps and a.outputs
        assert a.skill_path.exists(), f"{a.slug} has no SKILL.md"


def test_mandates_only_reference_real_agents():
    known = set(catalog.agents())
    for m in catalog.mandates().values():
        assert m.stages
        for _, slugs in m.stages:
            assert slugs
            assert set(slugs) <= known, f"{m.id} references an unknown agent"


def test_full_desk_covers_every_agent():
    assert set(catalog.mandates()["full"].agents) == set(catalog.agents())


def test_no_agent_repeats_within_a_mandate():
    for m in catalog.mandates().values():
        assert len(m.agents) == len(set(m.agents)), f"{m.id} runs an agent twice"


def test_categories_partition_the_roster():
    covered = [s for c in catalog.categories() for s in c.slugs]
    assert sorted(covered) == sorted(catalog.agents())
    assert len(covered) == len(set(covered))


def test_search_finds_by_method_language():
    assert "dcf-modeling" in {a.slug for a in catalog.search("WACC")}
    assert "negotiation-support" in {a.slug for a in catalog.search("BATNA")}
    assert catalog.search("zzzznope") == []
    assert len(catalog.search("")) == 42


# ── prompts ──────────────────────────────────────────────────────────────

def test_system_prompt_carries_the_whole_method():
    agent = catalog.get("dcf-modeling")
    prompt = catalog.system_prompt(agent)
    for step in agent.steps:
        assert step.head in prompt
    for out in agent.outputs:
        assert out in prompt


def test_deal_file_is_declared_authoritative():
    agent = catalog.get("dcf-modeling")
    prompt = catalog.user_prompt(
        agent,
        mandate=catalog.mandates()["valuation"],
        subject="ADNOC Drilling",
        deal_file="ENTITY: ADNOC Drilling\nRevenue AED 15.2bn",
    )
    assert "Authoritative" in prompt
    assert "AED 15.2bn" in prompt
    assert "deal file wins" in catalog.system_prompt(agent)


def test_empty_brief_tells_the_agent_to_flag_gaps():
    prompt = catalog.user_prompt(
        catalog.get("dcf-modeling"),
        mandate=catalog.mandates()["valuation"],
        subject="X",
    )
    assert "flag every gap" in prompt


def test_acronyms_survive_the_task_line():
    prompt = catalog.user_prompt(
        catalog.get("dcf-modeling"),
        mandate=catalog.mandates()["valuation"], subject="X",
    )
    assert "DCF + WACC build" in prompt


# ── routing ──────────────────────────────────────────────────────────────

def test_research_and_analysis_agents_route_apart():
    cfg = RunConfig(subject="X", mandate_id="valuation",
                    analysis_model="opus-4.8", research_model="sonar-pro")
    assert cfg.model_for(catalog.get("comparable-analysis")) == "sonar-pro"
    assert cfg.model_for(catalog.get("dcf-modeling")) == "opus-4.8"


def test_claude_research_agents_get_web_search(stub):
    run = Run(config=RunConfig(subject="X", mandate_id="valuation",
                               research_model="opus-4.8", web_search=True))
    _drain(run)
    expected = sum(1 for s in run.mandate.agents if catalog.get(s).researches)
    assert sum(1 for c in stub if c["web_search"]) == expected


def test_perplexity_agents_are_not_sent_a_web_tool(stub):
    """Sonar always searches; asking for the tool as well would be an error."""
    run = Run(config=RunConfig(subject="X", mandate_id="valuation",
                               research_model="sonar-pro", web_search=True))
    _drain(run)
    assert not any(c["web_search"] for c in stub)


def test_web_toggle_off_means_no_search(stub):
    run = Run(config=RunConfig(subject="X", mandate_id="valuation",
                               research_model="opus-4.8", web_search=False))
    _drain(run)
    assert not any(c["web_search"] for c in stub)


# ── orchestration ────────────────────────────────────────────────────────

def test_a_full_mandate_delivers_every_agent(stub):
    run = Run(config=RunConfig(subject="ADNOC Drilling", mandate_id="valuation"))
    _drain(run)
    done, total = run.progress()
    assert done == total == len(run.mandate.agents)
    assert len(stub) == total
    assert run.cost_tokens() == total * 1050


def test_later_agents_see_earlier_deliverables(stub):
    run = Run(config=RunConfig(subject="X", mandate_id="valuation"))
    _drain(run)
    last = run.mandate.stages[-1][1][0]
    assert len(run.upstream(last)) >= 8


def test_one_failure_does_not_stop_the_run(monkeypatch):
    def flaky(model_key, system, content, **kw):
        if "DCF Modeling agent" in system:
            raise providers.ProviderError("HTTP 429 — rate limited")
        model = providers.BY_KEY[model_key]
        return providers.Reply("fine", model.label, model.vendor)

    monkeypatch.setattr(orchestrator.providers, "complete", flaky)
    run = Run(config=RunConfig(subject="X", mandate_id="valuation"))
    _drain(run)

    assert run.status("dcf-modeling") == "failed"
    assert "429" in run.results["dcf-modeling"].error
    assert len(run.delivered) >= 9
    assert "Did not complete" in orchestrator.memo(run)


def test_retry_replaces_a_failure(stub, monkeypatch):
    run = Run(config=RunConfig(subject="X", mandate_id="valuation"))
    run.results["dcf-modeling"] = orchestrator.Result(
        slug="dcf-modeling", error="boom")
    assert run.status("dcf-modeling") == "failed"

    result = orchestrator.retry(run, "dcf-modeling")
    assert result.ok
    assert run.status("dcf-modeling") == "done"
    assert run.order.count("dcf-modeling") == 1


def test_memo_is_ordered_and_complete(stub):
    run = Run(config=RunConfig(subject="ADNOC Drilling", mandate_id="valuation",
                               deal_file="ENTITY: ADNOC Drilling"))
    _drain(run)
    memo = orchestrator.memo(run, ["FY25.pdf"])

    assert memo.startswith("# Valuation sprint — ADNOC Drilling")
    assert "## Deal file" in memo
    assert "FY25.pdf" in memo
    for label, _ in run.mandate.stages:
        assert f"## {label}" in memo
    positions = [memo.index(f"## {lbl}") for lbl, _ in run.mandate.stages]
    assert positions == sorted(positions), "stages are out of order"


# ── intake ───────────────────────────────────────────────────────────────

def test_csv_reads_as_text():
    src = intake.read("f.csv", b"Period,Revenue\nFY25,15200\n")
    assert src.kind == "sheet" and src.ready and "15200" in src.text


def test_xlsx_flattens_every_sheet():
    from openpyxl import Workbook

    wb = Workbook()
    wb.active.title = "IS"
    wb.active.append(["Period", "Revenue"])
    wb.active.append(["FY25", 15200])
    wb.create_sheet("BS").append(["Net debt", 4200])
    buf = io.BytesIO()
    wb.save(buf)

    src = intake.read("m.xlsx", buf.getvalue())
    assert "sheet: IS" in src.text and "sheet: BS" in src.text
    assert "15200" in src.text and "4200" in src.text


def test_docx_picks_up_tables_as_well_as_paragraphs():
    import docx

    doc = docx.Document()
    doc.add_paragraph("Net debt AED 4.2bn.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Shares"
    table.rows[0].cells[1].text = "16.0bn"
    buf = io.BytesIO()
    doc.save(buf)

    src = intake.read("m.docx", buf.getvalue())
    assert "4.2bn" in src.text and "16.0bn" in src.text


def test_images_go_up_as_native_blocks():
    src = intake.read("shot.png", b"\x89PNG\r\n\x1a\n" + b"0" * 40)
    assert src.blocks and src.blocks[0]["type"] == "image"
    assert src.blocks[0]["source"]["media_type"] == "image/png"


def test_unsupported_and_oversized_are_reported_not_raised():
    assert "unsupported" in intake.read("a.zip", b"x").error
    assert "over" in intake.read("big.pdf", b"0" * (intake.MAX_BYTES + 1)).error


def test_entity_is_lifted_from_the_digest():
    src = intake.Source("x", 1, "pdf")
    src.digest = "ENTITY: ADNOC Drilling PJSC\nPERIOD: FY25"
    assert src.entity == "ADNOC Drilling PJSC"
    src.digest = "ENTITY: unclear\nPERIOD: n/a"
    assert src.entity == ""


def test_digest_attaches_a_fact_base(stub):
    src = intake.digest(intake.read("f.csv", b"Revenue,15200\n"), "haiku-4.5")
    assert src.digest and not src.error
    assert stub[0]["model"] == "haiku-4.5"
    assert stub[0]["web_search"] is False


def test_deal_file_only_includes_read_sources():
    a = intake.Source("a.pdf", 1, "pdf")
    a.digest = "ENTITY: A"
    b = intake.Source("b.png", 1, "image")
    combined = intake.deal_file([a, b])
    assert "SOURCE: a.pdf" in combined
    assert "b.png" not in combined


# ── chat ─────────────────────────────────────────────────────────────────

def test_chat_context_carries_deal_file_and_deliverables(stub):
    run = Run(config=RunConfig(subject="ADNOC Drilling", mandate_id="valuation"))
    _drain(run)
    ctx = chat.context(run, "ENTITY: ADNOC Drilling\nRevenue AED 15.2bn", "ADNOC Drilling")

    assert "AED 15.2bn" in ctx
    assert "DCF Modeling" in ctx
    assert "agents have reported" in ctx


def test_chat_survives_an_empty_floor():
    ctx = chat.context(None, "", "")
    assert "AGENTS AVAILABLE" in ctx


def test_chat_sends_history_and_attachments(stub):
    history = [chat.Turn("user", "hello"), chat.Turn("assistant", "hi")]
    reply = chat.ask(
        "What is the EV/EBITDA?",
        model_key="opus-4.8", history=history, run=None,
        deal_file="EV/EBITDA 6.4x", subject="X",
        attachments=[{"type": "text", "text": "<attached>data</attached>"}],
    )
    assert reply.text
    content = stub[0]["content"]
    assert content[-1]["text"] == "What is the EV/EBITDA?"
    assert any("attached" in str(b) for b in content)


# ── providers ────────────────────────────────────────────────────────────

def test_model_table_is_coherent():
    assert len(providers.MODELS) == len(providers.BY_KEY)
    assert providers.ANTHROPIC_MODELS and providers.PERPLEXITY_MODELS
    for m in providers.PERPLEXITY_MODELS:
        assert m.searches, f"{m.key} should search natively"
    for m in providers.ANTHROPIC_MODELS:
        assert not m.searches, f"{m.key} searches only when given the tool"


def test_missing_key_is_a_clean_error(monkeypatch):
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    with pytest.raises(providers.ProviderError, match="ANTHROPIC_API_KEY"):
        providers.complete("opus-4.8", "s", "u")


def test_unknown_model_is_rejected():
    with pytest.raises(providers.ProviderError, match="Unknown model"):
        providers.complete("gpt-9", "s", "u")


# ── helper ───────────────────────────────────────────────────────────────

def _drain(run: Run) -> None:
    while not run.finished:
        orchestrator.run_stage(run, run.stage)
        run.stage += 1
