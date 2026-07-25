"""Turning attached documents into the deal file.

Each source is read once into a structured fact base. The agents then work
from that fact base rather than from raw documents, which keeps 42 agent calls
from each re-reading a 200-page annual report.
"""

from __future__ import annotations

import base64
import io
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from . import providers

MAX_BYTES = 25 * 1024 * 1024
TEXT_EXT = {".txt", ".md", ".markdown", ".json", ".log", ".text", ".rst"}
SHEET_EXT = {".xlsx", ".xlsm", ".xltx", ".csv", ".tsv"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".gif", ".webp"}
MEDIA = {
    ".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg",
    ".gif": "image/gif", ".webp": "image/webp",
}

INTAKE_SYSTEM = """You are the intake analyst on a deal team. You read one \
source document and write the fact base the rest of the team will work from.

Open with these two lines:
ENTITY: <the company or asset this document concerns, or "unclear">
PERIOD: <the periods the document covers>

Then short labelled sections, including only those the document actually supports:
- Income statement — revenue, gross profit, EBITDA, EBIT, net income, by period
- Balance sheet and capital structure — cash, gross and net debt, equity, share \
count, facilities and maturities
- Cash flow — operating cash flow, capex, free cash flow, dividends, buybacks
- Segments and operating KPIs — whatever this business actually reports
- Guidance and outlook — forward-looking statements, attributed as the \
company's own guidance
- One-offs, disputes, and risks

RULES
- Report only what appears in this document. Never estimate, never fill a gap, \
never supply a figure from general knowledge.
- Attach currency, units, and period to every figure.
- Omit any section the document does not support. Do not write "not disclosed" lines.
- No summary paragraph, no commentary, no recommendation. This is a fact base, \
not analysis.
- Under 450 words."""


@dataclass
class Source:
    name: str
    size: int
    kind: str                       # text | sheet | pdf | doc | image
    text: str = ""                  # extracted text, when extractable
    blocks: list[dict[str, Any]] = field(default_factory=list)  # native blocks
    digest: str = ""                # the fact base
    error: str = ""

    @property
    def ready(self) -> bool:
        return bool(self.text or self.blocks) and not self.error

    @property
    def entity(self) -> str:
        for line in self.digest.splitlines():
            if line.upper().startswith("ENTITY:"):
                value = line.split(":", 1)[1].strip().strip("*_ ")
                return "" if "unclear" in value.lower() else value
        return ""


def classify(name: str) -> str:
    ext = Path(name).suffix.lower()
    if ext == ".pdf":
        return "pdf"
    if ext == ".docx":
        return "doc"
    if ext in IMAGE_EXT:
        return "image"
    if ext in SHEET_EXT:
        return "sheet"
    if ext in TEXT_EXT:
        return "text"
    return "other"


def read(name: str, raw: bytes) -> Source:
    """Extract what we can locally. Scanned PDFs and images go up as-is."""
    kind = classify(name)
    src = Source(name=name, size=len(raw), kind=kind)

    if len(raw) > MAX_BYTES:
        src.error = f"over {MAX_BYTES // 1024 // 1024} MB — split it first"
        return src

    try:
        if kind == "text":
            src.text = raw.decode("utf-8", errors="replace")
        elif kind == "sheet":
            src.text = _sheet(name, raw)
        elif kind == "doc":
            src.text = _docx(raw)
        elif kind == "pdf":
            text = _pdf_text(raw)
            # A page of a text PDF yields ~1500 chars. Far less means it is
            # scanned, so hand the file to the model instead of guessing.
            if len(text.strip()) > 400:
                src.text = text
            else:
                src.blocks = [{
                    "type": "document",
                    "source": {
                        "type": "base64",
                        "media_type": "application/pdf",
                        "data": base64.standard_b64encode(raw).decode(),
                    },
                }]
        elif kind == "image":
            src.blocks = [{
                "type": "image",
                "source": {
                    "type": "base64",
                    "media_type": MEDIA.get(Path(name).suffix.lower(), "image/png"),
                    "data": base64.standard_b64encode(raw).decode(),
                },
            }]
        else:
            src.error = "unsupported type — use PDF, XLSX, CSV, DOCX, TXT or an image"
    except Exception as exc:  # noqa: BLE001 — reported per-file, run continues
        src.error = str(exc)
    return src


def _sheet(name: str, raw: bytes) -> str:
    if Path(name).suffix.lower() in {".csv", ".tsv"}:
        return raw.decode("utf-8", errors="replace")

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    out = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            if row is None or all(c is None for c in row):
                continue
            rows.append(",".join("" if c is None else str(c) for c in row))
            if len(rows) >= 400:
                rows.append("[…truncated]")
                break
        if rows:
            out.append(f"--- sheet: {ws.title} ---\n" + "\n".join(rows))
    wb.close()
    return "\n\n".join(out)


def _docx(raw: bytes) -> str:
    import docx

    doc = docx.Document(io.BytesIO(raw))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _pdf_text(raw: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(raw))
    pages = []
    for i, page in enumerate(reader.pages[:150]):
        try:
            pages.append(page.extract_text() or "")
        except Exception:  # noqa: BLE001 — one bad page shouldn't kill the file
            continue
    return "\n\n".join(pages)


def digest(src: Source, model_key: str) -> Source:
    """Send one source for extraction and attach the resulting fact base."""
    if not src.ready:
        return src

    if src.blocks:
        content = [
            *src.blocks,
            {"type": "text",
             "text": f"Source file: {src.name}\n\nExtract the fact base."},
        ]
    else:
        body = src.text if len(src.text) <= 120_000 else src.text[:120_000] + "\n[…truncated]"
        content = [{
            "type": "text",
            "text": (f"Source file: {src.name}\n\n<document>\n{body}\n</document>\n\n"
                     "Extract the fact base."),
        }]

    try:
        reply = providers.complete(
            model_key, INTAKE_SYSTEM, content, max_tokens=1200, web_search=False
        )
        src.digest = reply.text
        src.error = ""
    except providers.ProviderError as exc:
        src.error = str(exc)
    return src


def deal_file(sources: list[Source]) -> str:
    """Every fact base, concatenated, with its source named."""
    return "\n\n".join(
        f"SOURCE: {s.name}\n{s.digest}" for s in sources if s.digest
    )
